"""基于第46届觉观业务表，生成第47届觉观新版课程工作簿。"""
from __future__ import annotations

import argparse
import copy
from datetime import date, datetime, timedelta
from pathlib import Path
import re
import sys
import time
from typing import Any

from openpyxl import load_workbook
from sqlmodel import Session, select

REPO_ROOT = Path(__file__).resolve().parents[1]
if str(REPO_ROOT) not in sys.path:
    sys.path.insert(0, str(REPO_ROOT))

from backend.core.nianzhu_course_sheets import (  # noqa: E402
    CLOCKIN_CONFIG_COLUMNS,
    CLOCKIN_CONFIG_SHEET_KEY,
    CLOCKIN_DATA_COLUMNS,
    CLOCKIN_DATA_SHEET_KEY,
    VIDEO_CONFIG_COLUMNS,
    VIDEO_CONFIG_SHEET_KEY,
    VIDEO_DATA_COLUMNS,
    VIDEO_DATA_SHEET_KEY,
    video_config_url_from_lesson_id2,
)
from backend.core.note_sheet_access import ensure_attendance_sheet_anonymous_viewer  # noqa: E402
from backend.core.sheet_identity import allocate_new_sheet_identity, allocate_new_workbook_identity  # noqa: E402
from backend.core.sheet_refs import sheet_public_id, sheet_ref_aliases, workbook_public_id, workbook_ref_aliases  # noqa: E402
from backend.db import engine  # noqa: E402
from backend.models import SheetDocument, WorkbookDocument, WorkbookSheetLink  # noqa: E402
from scripts.import_legacy_attendance_workbook import (  # noqa: E402
    REGISTRATION_ACTION_LABELS,
    REGISTRATION_ACTION_ROW_NOTES,
    _attendance_document,
    _default_view_settings,
    _registration_column_configs,
)


DEFAULT_TEMPLATE_PATH = Path(r"C:\Users\kzche\Downloads\20260501第46届觉观.xlsx")
COURSE_NAME = "d260601第47届觉观"
WORKBOOK_TITLE = "第47届觉观"
SUMMARY_ONLINE_SHEET_NAME = "第47届觉观"
OWNER_KEY = "20260601-jueguan-47"
TEMPLATE_COURSE_START_DATE = date(2026, 5, 1)
COURSE_START_DATE = date(2026, 6, 1)
OFFICIAL_LESSON_COUNT = 21
STANDARD_REGISTRATION_COLUMNS = [
    "分组",
    "序号",
    "备注",
    "提交时间",
    "姓名",
    "微信昵称",
    "手机号",
    "错误手机号",
    "微信支付订单号",
    "订单日期",
    "商户订单号",
    "订单金额",
    "用户ID",
    "匹配得分",
    "参考信息",
    "关联用户ID",
    "报名项目（必填）",
    "性别（必填）",
    "微信昵称（必填）",
    "出生年月（必填）",
    "常住地址（必填）",
    "学历（必填）",
    "专业（必填）",
    "职业（必填）",
    "您的首要学习动机是？（必填）",
    "禅修经历（必填）",
    "在本系统开始学修的时间（必填）",
    "主修法门（必填）",
    "系统内觉观禅地面营经历（必填）",
    "系统外学修经历",
    "您了解我们课程的渠道是？（必填）",
    "版权承诺（必填）",
    "进入学习群",
    "提交者（自动）",
]
REGISTRATION_EXTRA_ACTION_LABELS = {
    "提交时间": ("registration_add_student", "新增学员"),
    "匹配得分": ("registration_composite_update", "综合更新"),
}

COURSE_LESSON_CONFIG = [
    {"lesson_id2": "l_69536897e4b0694ca162e9bd", "shop_id": 1, "start_date": "2026-06-01 05:20:00", "next_update": "2026-06-01 06:18:30", "video_duration": 3510},
    {"lesson_id2": "l_6953689ae4b0694c5b6c1057", "shop_id": 1, "start_date": "2026-06-02 05:20:00", "next_update": "2026-06-02 06:02:18", "video_duration": 2538},
    {"lesson_id2": "l_6953689be4b0694ca162e9c5", "shop_id": 1, "start_date": "2026-06-03 05:20:00", "next_update": "2026-06-03 06:10:31", "video_duration": 3031},
    {"lesson_id2": "l_6953689de4b0694c5b6c105c", "shop_id": 1, "start_date": "2026-06-04 05:20:00", "next_update": "2026-06-04 06:07:56", "video_duration": 2876},
    {"lesson_id2": "l_6953689fe4b0694ca162e9cc", "shop_id": 1, "start_date": "2026-06-05 05:20:00", "next_update": "2026-06-05 06:14:25", "video_duration": 3265},
    {"lesson_id2": "l_695368a1e4b0694ca162e9d0", "shop_id": 1, "start_date": "2026-06-06 05:20:00", "next_update": "2026-06-06 06:05:38", "video_duration": 2738},
    {"lesson_id2": "l_695368a3e4b0694c5b6c1066", "shop_id": 1, "start_date": "2026-06-07 05:20:00", "next_update": "2026-06-07 06:05:01", "video_duration": 2701},
    {"lesson_id2": "l_695368a5e4b0694c5b6c106e", "shop_id": 1, "start_date": "2026-06-08 05:20:00", "next_update": "2026-06-08 06:16:15", "video_duration": 3375},
    {"lesson_id2": "l_695368a7e4b0694ca162e9d6", "shop_id": 1, "start_date": "2026-06-09 05:20:00", "next_update": "2026-06-09 06:06:06", "video_duration": 2766},
    {"lesson_id2": "l_695368a8e4b0694c5b6c1072", "shop_id": 1, "start_date": "2026-06-10 05:20:00", "next_update": "2026-06-10 06:11:41", "video_duration": 3101},
    {"lesson_id2": "l_695368aae4b0694ca162e9df", "shop_id": 1, "start_date": "2026-06-11 05:20:00", "next_update": "2026-06-11 06:17:27", "video_duration": 3447},
    {"lesson_id2": "l_695368ace4b0694ca162e9e3", "shop_id": 1, "start_date": "2026-06-12 05:20:00", "next_update": "2026-06-12 06:13:27", "video_duration": 3207},
    {"lesson_id2": "l_695368aee4b0694c5b6c107b", "shop_id": 1, "start_date": "2026-06-13 05:20:00", "next_update": "2026-06-13 06:15:16", "video_duration": 3316},
    {"lesson_id2": "l_695368b0e4b0694ca162e9e9", "shop_id": 1, "start_date": "2026-06-14 05:20:00", "next_update": "2026-06-14 05:59:05", "video_duration": 2345},
    {"lesson_id2": "l_695368b2e4b0694ca162e9f3", "shop_id": 1, "start_date": "2026-06-15 05:20:00", "next_update": "2026-06-15 05:59:07", "video_duration": 2347},
    {"lesson_id2": "l_695368b4e4b0694ca162ea01", "shop_id": 1, "start_date": "2026-06-16 05:20:00", "next_update": "2026-06-16 06:25:14", "video_duration": 3914},
    {"lesson_id2": "l_695368b5e4b0694c5b6c108d", "shop_id": 1, "start_date": "2026-06-17 05:20:00", "next_update": "2026-06-17 06:34:11", "video_duration": 4451},
    {"lesson_id2": "l_695368b7e4b0694ca162ea0b", "shop_id": 1, "start_date": "2026-06-18 05:20:00", "next_update": "2026-06-18 06:20:35", "video_duration": 3635},
    {"lesson_id2": "l_695368b9e4b0694ca162ea0d", "shop_id": 1, "start_date": "2026-06-19 05:20:00", "next_update": "2026-06-19 06:45:58", "video_duration": 5158},
    {"lesson_id2": "l_695368bbe4b0694c5b6c10a3", "shop_id": 1, "start_date": "2026-06-20 05:20:00", "next_update": "2026-06-20 06:25:20", "video_duration": 3920},
    {"lesson_id2": "l_695368bce4b0694ca162ea11", "shop_id": 1, "start_date": "2026-06-21 05:20:00", "next_update": "2026-06-21 06:48:21", "video_duration": 5301},
]
CLOCKIN_CONFIG_URL = "https://admin.xiaoe-tech.com/t/clock_admin/index#/punchDetail/diaryList?activity_id=ac_695377129416a_ltbHq1FH&markType=&miniMiddleUrl=https%3A%2F%2Fapporrfwkpb5562.h5.xet.citv.cn%2Fxiaoe_clock%2Fmini_middle%3Factivity_id%3Dac_695377129416a_ltbHq1FH%26app_id%3Dapporrfwkpb5562"


def _text(value: Any) -> str:
    return str(value or "").strip()


def _shift_day(month: int, day: int) -> date | None:
    try:
        source_day = date(TEMPLATE_COURSE_START_DATE.year, month, day)
    except ValueError:
        return None
    return COURSE_START_DATE + (source_day - TEMPLATE_COURSE_START_DATE)


def _fmt_day(day: date) -> str:
    return f"{day.month}月{day.day}日"


def _shift_template_dates(text: str) -> str:
    def repl_chinese(match: re.Match[str]) -> str:
        shifted = _shift_day(5, int(match.group(1)))
        return _fmt_day(shifted) if shifted else match.group(0)

    def repl_iso(match: re.Match[str]) -> str:
        shifted = _shift_day(5, int(match.group(2)))
        return _fmt_day(shifted) if shifted else match.group(0)

    anchor = COURSE_START_DATE - timedelta(days=1)
    text = re.sub(r"5月(\d{1,2})日", repl_chinese, text)
    text = re.sub(r"2026([/-])0?5\1(\d{1,2})", repl_iso, text)
    text = re.sub(r"2026([/-])0?4[/-]30", f"{anchor.year}-{'%02d' % anchor.month}-{'%02d' % anchor.day}", text)
    return text.replace("第46届觉观", "第47届觉观").replace("第46届", "第47届")


def _map_text(value: Any) -> Any:
    if isinstance(value, list):
        return [_map_text(item) for item in value]
    if isinstance(value, dict):
        return {key: _map_text(item) for key, item in value.items()}
    if not isinstance(value, str):
        return value
    if "最近运行更新时间" in value:
        return "最近运行更新时间：\n待首次同步"
    return _shift_template_dates(value)


def _remove_column(document: dict[str, Any], column_index: int) -> None:
    columns = list(document.get("columns") or [])
    if not 0 <= column_index < len(columns):
        return
    removed = columns[column_index]
    document["columns"] = columns[:column_index] + columns[column_index + 1:]
    for key in ("rows", "grid_rows"):
        document[key] = [
            list(row)[:column_index] + list(row)[column_index + 1:]
            for row in document.get(key) or []
            if isinstance(row, list)
        ]
    widths = list(document.get("column_widths") or [])
    if len(widths) > column_index:
        document["column_widths"] = widths[:column_index] + widths[column_index + 1:]
    configs = dict(document.get("column_configs") or {})
    configs.pop(removed, None)
    document["column_configs"] = configs

    adjusted_meta: dict[str, Any] = {}
    for key, value in dict(document.get("cell_meta") or {}).items():
        try:
            row_text, col_text = str(key).split(":", 1)
            col = int(col_text)
        except ValueError:
            adjusted_meta[key] = value
            continue
        if col == column_index:
            continue
        adjusted_meta[f"{row_text}:{col - 1 if col > column_index else col}"] = value
    document["cell_meta"] = adjusted_meta

    merged_cells = []
    for item in document.get("merged_cells") or []:
        col = int(item.get("col") or 0)
        colspan = int(item.get("colspan") or 1)
        if col <= column_index < col + colspan:
            colspan -= 1
            if colspan <= 0:
                continue
        elif col > column_index:
            col -= 1
        merged_cells.append({**item, "col": col, "colspan": colspan})
    document["merged_cells"] = merged_cells


def _set_cell_style(document: dict[str, Any], row: int, col: int, style: dict[str, Any]) -> None:
    cell_meta = dict(document.get("cell_meta") or {})
    key = f"{row}:{col}"
    meta = dict(cell_meta.get(key) or {})
    next_style = dict(meta.get("style") or {})
    next_style.update(style)
    meta["style"] = next_style
    cell_meta[key] = meta
    document["cell_meta"] = cell_meta


def _set_cell_link(document: dict[str, Any], row: int, col: int, url: str | None) -> None:
    cell_meta = dict(document.get("cell_meta") or {})
    key = f"{row}:{col}"
    meta = dict(cell_meta.get(key) or {})
    if url:
        meta["link"] = {"url": url}
    else:
        meta.pop("link", None)
    cell_meta[key] = meta
    document["cell_meta"] = cell_meta


def _attendance_defined_names() -> list[dict[str, str]]:
    anchor = COURSE_START_DATE - timedelta(days=1)
    return [
        {"name": "返款周期", "formula": f"=INT(TODAY()-DATE({anchor.year},{anchor.month},{anchor.day}))", "comment": "当前课程返款周期天数，按开课前一天到今天计算，支持开课前负数。", "scope": "worksheet"},
        {"name": "返款说明", "formula": '="第"&返款周期&"天返款"', "comment": "返款操作说明文本。", "scope": "worksheet"},
        {"name": "返款ID后缀", "formula": '="_dayc"&返款周期', "comment": "觉观返款批次 ID 后缀。", "scope": "worksheet"},
    ]


def _normalize_refund_layout(document: dict[str, Any]) -> None:
    columns = list(document.get("columns") or [])
    refund_start = columns.index("完成视频数")
    refund_end = columns.index("总应返款")
    operation_start = columns.index("订单金额")
    operation_end = columns.index("当前应返款")
    clockin_index = columns.index("打卡数")
    grid_rows = [list(row) for row in document.get("grid_rows") or [] if isinstance(row, list)]
    if grid_rows:
        for col in range(refund_start, min(clockin_index + 1, len(grid_rows[0]))):
            grid_rows[0][col] = ""
        grid_rows[0][refund_start] = "返款总计（数据仅在每天早上更新一次，不是实时更新！）"
        grid_rows[0][operation_start] = "返款操作"
        grid_rows[0][clockin_index] = "打卡"
        document["grid_rows"] = grid_rows
    merged = []
    for item in document.get("merged_cells") or []:
        is_refund_header = int(item.get("row") or 0) == 0 and int(item.get("col") or 0) < clockin_index + 1
        if is_refund_header and int(item.get("col") or 0) + int(item.get("colspan") or 1) > refund_start:
            continue
        merged.append(item)
    merged.append({"row": 0, "col": refund_start, "rowspan": 1, "colspan": refund_end - refund_start + 1})
    merged.append({"row": 0, "col": operation_start, "rowspan": 1, "colspan": operation_end - operation_start + 1})
    document["merged_cells"] = merged


def _lesson_label(config: dict[str, Any], number: int) -> str:
    start_dt = datetime.strptime(config["start_date"], "%Y-%m-%d %H:%M:%S")
    end_dt = datetime.strptime(config["next_update"], "%Y-%m-%d %H:%M:%S")
    return f"{start_dt:%H:%M}~{end_dt:%H:%M} 第{number:02d}课"


def _adapt_attendance_document(document: dict[str, Any]) -> dict[str, Any]:
    doc = _map_text(copy.deepcopy(document))
    for column_name in ("返款配置", "已返款"):
        columns = list(doc.get("columns") or [])
        if column_name in columns:
            _remove_column(doc, columns.index(column_name))

    columns = list(doc.get("columns") or [])
    grid_rows = [list(row) for row in doc.get("grid_rows") or [] if isinstance(row, list)]
    first_lesson_index = next(index for index, column in enumerate(columns) if re.search(r"第\s*0*1\s*课", column))
    for index, config in enumerate(COURSE_LESSON_CONFIG, start=1):
        col = first_lesson_index + index - 1
        start_day = COURSE_START_DATE + timedelta(days=index - 1)
        label = _lesson_label(config, index)
        columns[col] = label
        grid_rows[0][col] = f"{_fmt_day(start_day)}~{_fmt_day(start_day + timedelta(days=4))}"
        grid_rows[1][col] = label
        _set_cell_link(doc, 1, col, f"https://admin.xiaoe-tech.com/t/live_management#/userOperation?id={config['lesson_id2']}&tabName=UserManage")

    extra_start = first_lesson_index + OFFICIAL_LESSON_COUNT
    if len(columns) > extra_start:
        columns[extra_start] = "20:00~22:00 觉观同学会"
        grid_rows[0][extra_start] = "6月22日"
        grid_rows[1][extra_start] = "20:00~22:00 觉观同学会"
        _set_cell_link(doc, 1, extra_start, None)
    if len(columns) > extra_start + 1:
        columns[extra_start + 1] = "20:00~22:00 优秀学员学修分享会"
        grid_rows[0][extra_start + 1] = "6月23日"
        grid_rows[1][extra_start + 1] = "20:00~22:00 优秀学员学修分享会"
        _set_cell_link(doc, 1, extra_start + 1, None)

    feedback_index = columns.index("打卡应返款")
    order_amount_index = columns.index("总应返款")
    refund_period_index = columns.index("订单金额")
    status_index = columns.index("当前应返款")
    clockin_index = columns.index("打卡数")
    grid_rows[2][feedback_index] = "点击我-反馈考勤返款数据问题。义工有空会统一处理。"
    grid_rows[2][order_amount_index] = 499
    grid_rows[2][refund_period_index] = '="第"&返款周期&"天"'
    grid_rows[2][status_index] = "最近运行更新时间：\n待首次同步"
    grid_rows[2][clockin_index] = "只统计正课的打卡次数！！！"
    doc["columns"] = columns
    doc["grid_rows"] = grid_rows[: int(doc.get("data_start_row") or 0)]
    doc["rows"] = []
    data_start = int(doc.get("data_start_row") or 0)
    doc["cell_meta"] = {
        key: value
        for key, value in dict(doc.get("cell_meta") or {}).items()
        if int(str(key).split(":", 1)[0]) < data_start
    }
    for index, config in enumerate(COURSE_LESSON_CONFIG, start=1):
        _set_cell_link(doc, 1, first_lesson_index + index - 1, video_config_url_from_lesson_id2(config["lesson_id2"]))
    _set_cell_link(doc, 1, clockin_index, CLOCKIN_CONFIG_URL)

    _normalize_refund_layout(doc)
    _set_cell_style(doc, 2, 1, {"font_size": 16, "text_color": "#FF0000", "text_align": "center", "vertical_align": "middle", "background_color": "#D9D9D9"})
    _set_cell_style(doc, 2, feedback_index, {"background_color": "#D9D9D9", "text_color": "#333333", "text_align": "center", "vertical_align": "middle"})
    _set_cell_style(doc, 2, order_amount_index, {"background_color": "#D9D9D9", "text_color": "#555555", "text_align": "center", "vertical_align": "middle"})
    _set_cell_style(doc, 2, refund_period_index, {"background_color": "#D9D9D9", "text_color": "#555555", "text_align": "center", "vertical_align": "middle"})
    _set_cell_style(doc, 2, status_index, {"background_color": "#D9D9D9", "text_color": "#555555", "text_align": "center", "vertical_align": "middle"})
    _set_cell_style(doc, 2, first_lesson_index, {"text_align": "left", "vertical_align": "middle", "background_color": "#D9D9D9"})
    doc["defined_names"] = _attendance_defined_names()
    doc["source_meta"] = {
        **dict(doc.get("source_meta") or {}),
        "course_name": COURSE_NAME,
        "business_template": "20260501第46届觉观",
        "business_structure": "jueguan_46",
        "columns": len(columns),
    }
    return doc


def _standard_registration_document() -> dict[str, Any]:
    columns = list(STANDARD_REGISTRATION_COLUMNS)
    action_row = [""] * len(columns)
    for header, (_action_type, label) in REGISTRATION_ACTION_LABELS.items():
        if header in columns:
            action_row[columns.index(header)] = label
    for header, note in REGISTRATION_ACTION_ROW_NOTES.items():
        if header in columns:
            action_row[columns.index(header)] = note

    cell_meta: dict[str, Any] = {
        f"0:{column_index}": {"style": {"background_color": "#9DC3E6"}}
        for column_index in range(len(columns))
    }
    for header, (action_type, label) in {
        **REGISTRATION_ACTION_LABELS,
        **REGISTRATION_EXTRA_ACTION_LABELS,
    }.items():
        if header in columns:
            cell_meta[f"1:{columns.index(header)}"] = {"action": {"type": action_type, "label": label}}

    return {
        "schema_version": 1,
        "columns": columns,
        "rows": [],
        "grid_rows": [columns, action_row],
        "data_start_row": 2,
        "field_row_index": 0,
        "column_widths": [88, 72, 120, 150, 100, 120, 120, 120, 210, 100, 180, 90, 220, 90, 180, *([160] * (len(columns) - 15))],
        "column_configs": _registration_column_configs(columns, action_row),
        "cell_meta": cell_meta,
        "merged_cells": [],
        "header_groups": [],
        "formula_reference_origin": "sheet_v2",
        "view_settings": _default_view_settings(row_marker_numbering="global"),
        "source_meta": {
            "course_name": COURSE_NAME,
            "business_template": "20260601第41届念住",
            "business_structure": "standard_registration",
            "columns": len(columns),
        },
    }


def _simple_document(columns: list[str], rows: list[list[Any]], source_meta: dict[str, Any] | None = None) -> dict[str, Any]:
    return {
        "schema_version": 1,
        "columns": columns,
        "rows": rows,
        "grid_rows": [columns, *rows],
        "data_start_row": 1,
        "field_row_index": 0,
        "column_configs": {},
        "column_widths": [120] * len(columns),
        "cell_meta": {},
        "merged_cells": [],
        "header_groups": [],
        "formula_reference_origin": "sheet_v2",
        "view_settings": {"show_row_numbers": True, "row_marker_numbering": "page", "row_marker_origin": "sheet", "show_column_markers": True, "column_marker_style": "letters", "height_mode": "fill", "pagination": {"enabled": True, "page_size": 100}},
        "source_meta": dict(source_meta or {}),
    }


def _lesson_end_date(config: dict[str, Any]) -> str:
    return (datetime.strptime(config["next_update"], "%Y-%m-%d %H:%M:%S") + timedelta(days=5)).strftime("%Y-%m-%d %H:%M:%S")


def _storage_documents() -> dict[str, dict[str, Any]]:
    source_meta = {
        "course_name": COURSE_NAME,
        "business_template": "20260501第46届觉观",
        "business_structure": "jueguan_46",
        "video_refund_rule_mode": "timed_text",
        "timed_video_rules": {"当堂": 19, "第1天": 14, "第2天": 9, "第3天": 4, "回放": 0},
    }
    video_config_rows = [
        [
            index,
            config["start_date"],
            _lesson_end_date(config),
            config["next_update"],
            config["lesson_id2"],
            config["shop_id"],
            f"第{index:02d}课",
            config["video_duration"],
        ]
        for index, config in enumerate(COURSE_LESSON_CONFIG, start=1)
    ]
    return {
        VIDEO_CONFIG_SHEET_KEY: _simple_document(VIDEO_CONFIG_COLUMNS, video_config_rows, {**source_meta, "lesson_count": OFFICIAL_LESSON_COUNT}),
        VIDEO_DATA_SHEET_KEY: _simple_document(VIDEO_DATA_COLUMNS, [], source_meta),
        CLOCKIN_CONFIG_SHEET_KEY: _simple_document(CLOCKIN_CONFIG_COLUMNS, [[1, "打卡数", CLOCKIN_CONFIG_URL, "", "", "", "", ""]], source_meta),
        CLOCKIN_DATA_SHEET_KEY: _simple_document(CLOCKIN_DATA_COLUMNS, [], source_meta),
    }


def _load_documents(template_path: Path) -> tuple[dict[str, Any], dict[str, Any]]:
    workbook = load_workbook(template_path, read_only=False, data_only=False)
    try:
        attendance = _adapt_attendance_document(_attendance_document(workbook["考勤表"]))
        registration = _standard_registration_document()
        return attendance, registration
    finally:
        workbook.close()


def _find_or_create_workbook(session: Session, owner_user_id: int) -> WorkbookDocument:
    workbook = session.exec(select(WorkbookDocument).where(WorkbookDocument.title == WORKBOOK_TITLE)).first()
    if workbook is None:
        workbook = session.exec(select(WorkbookDocument).where(WorkbookDocument.title == COURSE_NAME)).first()
    if workbook is not None:
        return workbook
    identity = allocate_new_workbook_identity(session)
    now = time.time()
    workbook = WorkbookDocument(
        id=identity.primary_id,
        numeric_id=identity.numeric_id,
        legacy_id=identity.legacy_id,
        title=WORKBOOK_TITLE,
        owner_user_id=owner_user_id,
        created_by_user_id=owner_user_id,
        updated_by_user_id=owner_user_id,
        created_at=now,
        updated_at=now,
    )
    session.add(workbook)
    session.flush()
    return workbook


def _upsert_sheet(session: Session, workbook: WorkbookDocument, sheet_key: str, title: str, document: dict[str, Any]) -> SheetDocument:
    sheet = session.exec(
        select(SheetDocument)
        .where(SheetDocument.owner_type == "course_workbook")
        .where(SheetDocument.owner_key == OWNER_KEY)
        .where(SheetDocument.sheet_key == sheet_key)
    ).first()
    now = time.time()
    if sheet is None:
        identity = allocate_new_sheet_identity(session)
        sheet = SheetDocument(
            id=identity.primary_id,
            numeric_id=identity.numeric_id,
            legacy_id=identity.legacy_id,
            scope="notes",
            owner_type="course_workbook",
            owner_key=OWNER_KEY,
            sheet_key=sheet_key,
            title=title,
            engine="handsontable",
            document_json=document,
            version=1,
            owner_user_id=workbook.owner_user_id,
            created_by_user_id=workbook.created_by_user_id,
            updated_by_user_id=workbook.updated_by_user_id,
            created_at=now,
            updated_at=now,
        )
    else:
        sheet.title = title
        sheet.document_json = document
        sheet.version = max(int(sheet.version or 1), 1) + 1
        sheet.updated_by_user_id = workbook.updated_by_user_id or workbook.owner_user_id
        sheet.updated_at = now
    session.add(sheet)
    session.flush()
    return sheet


def _ensure_link(session: Session, workbook: WorkbookDocument, sheet: SheetDocument, order_index: int) -> None:
    workbook_ref = workbook_public_id(workbook)
    sheet_ref = sheet_public_id(sheet)
    link = session.exec(
        select(WorkbookSheetLink)
        .where(WorkbookSheetLink.workbook_id.in_(workbook_ref_aliases(workbook)))
        .where(WorkbookSheetLink.sheet_id.in_(sheet_ref_aliases(sheet)))
    ).first()
    if link is None:
        link = WorkbookSheetLink(workbook_id=workbook_ref, sheet_id=sheet_ref, order_index=order_index, created_at=time.time())
    else:
        link.workbook_id = workbook_ref
        link.sheet_id = sheet_ref
        link.order_index = order_index
    session.add(link)


def _link_summary_c4(session: Session, workbook: WorkbookDocument, attendance: SheetDocument) -> bool:
    summary = session.exec(select(SheetDocument).where(SheetDocument.numeric_id == 4)).first()
    if summary is None:
        return False
    document = copy.deepcopy(dict(summary.document_json or {}))
    columns = list(document.get("columns") or [])
    rows = [list(row) for row in document.get("rows") or [] if isinstance(row, list)]
    try:
        online_col = columns.index("在线考勤表")
        course_name_col = columns.index("课程名称")
    except ValueError:
        return False
    target_row = next(
        (
            index
            for index, row in enumerate(rows)
            if course_name_col < len(row) and _text(row[course_name_col]) == "第47届觉观"
        ),
        -1,
    )
    if target_row < 0:
        return False
    rows[target_row][online_col] = SUMMARY_ONLINE_SHEET_NAME
    cell_meta = dict(document.get("cell_meta") or {})
    meta = dict(cell_meta.get(f"{target_row}:{online_col}") or {})
    meta["link"] = {"url": f"/workbook/{int(workbook.numeric_id or 0)}?sheet={int(attendance.numeric_id or 0)}"}
    cell_meta[f"{target_row}:{online_col}"] = meta
    document["rows"] = rows
    document["cell_meta"] = cell_meta
    data_start = int(document.get("data_start_row") or 0)
    document["grid_rows"] = list(document.get("grid_rows") or [])[:data_start] + rows
    summary.document_json = document
    summary.version = max(int(summary.version or 1), 1) + 1
    summary.updated_by_user_id = workbook.updated_by_user_id or workbook.owner_user_id
    summary.updated_at = time.time()
    session.add(summary)
    return True


def run(*, apply: bool, template_path: Path, owner_user_id: int = 2) -> dict[str, Any]:
    attendance_document, registration_document = _load_documents(template_path)
    storage_documents = _storage_documents()
    with Session(engine) as session:
        workbook = _find_or_create_workbook(session, owner_user_id)
        workbook.title = WORKBOOK_TITLE
        workbook.updated_by_user_id = owner_user_id
        workbook.updated_at = time.time()
        session.add(workbook)

        specs = [
            ("attendance", "考勤表", 5, attendance_document),
            ("registration", "报名表", 10, registration_document),
            (VIDEO_CONFIG_SHEET_KEY, "视频配置", 30, storage_documents[VIDEO_CONFIG_SHEET_KEY]),
            (VIDEO_DATA_SHEET_KEY, "视频数据", 40, storage_documents[VIDEO_DATA_SHEET_KEY]),
            (CLOCKIN_CONFIG_SHEET_KEY, "打卡配置", 50, storage_documents[CLOCKIN_CONFIG_SHEET_KEY]),
            (CLOCKIN_DATA_SHEET_KEY, "打卡数据", 60, storage_documents[CLOCKIN_DATA_SHEET_KEY]),
        ]
        sheets = []
        for sheet_key, title, order_index, document in specs:
            sheet = _upsert_sheet(session, workbook, sheet_key, title, document)
            _ensure_link(session, workbook, sheet, order_index)
            ensure_attendance_sheet_anonymous_viewer(session, sheet)
            sheets.append(sheet)

        summary_linked = _link_summary_c4(session, workbook, sheets[0])
        if apply:
            session.commit()
        else:
            session.rollback()

        return {
            "mode": "APPLY" if apply else "DRY-RUN",
            "workbook_id": int(workbook.numeric_id or 0),
            "attendance_sheet_id": int(sheets[0].numeric_id or 0),
            "summary_c4_linked": summary_linked,
            "sheets": [{"sheet_key": sheet.sheet_key, "title": sheet.title, "id": int(sheet.numeric_id or 0), "rows": len((sheet.document_json or {}).get("rows") or [])} for sheet in sheets],
        }


def main() -> None:
    parser = argparse.ArgumentParser(description="重建第47届觉观为 46届业务结构 + 新版课程工作簿框架。")
    parser.add_argument("--apply", action="store_true", help="实际写入数据库；默认 dry-run。")
    parser.add_argument("--template", type=Path, default=DEFAULT_TEMPLATE_PATH)
    args = parser.parse_args()
    print(run(apply=args.apply, template_path=args.template))


if __name__ == "__main__":
    main()
