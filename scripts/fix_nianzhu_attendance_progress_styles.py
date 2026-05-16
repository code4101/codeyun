from __future__ import annotations

import argparse
import copy
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
import re
import sys
import time
from typing import Any

from openpyxl import load_workbook
from sqlmodel import Session, select

REPO_ROOT = Path(__file__).resolve().parents[1]
if str(REPO_ROOT) not in sys.path:
    sys.path.insert(0, str(REPO_ROOT))

from backend.api.note_sheets import _extract_document_rows, _normalize_document_columns, _normalize_document_data_start_row
from backend.db import engine
from backend.models import SheetDocument
from scripts.apply_nianzhu_attendance_progress_styles import apply_nianzhu_course_progress_styles_to_document
from scripts.import_legacy_attendance_workbook import _attendance_document


DEFAULT_WORKBOOK_PATH = Path("C:/Users/kzche/Downloads/20250106念住闯关.xlsx")
ATTENDANCE_SHEET_NUMERIC_ID = 21
SOURCE_SHEET_TITLES = (
    "考勤表",
    "考勤表2025年5月22日以后",
    "考勤表2025年5月22日以前",
)

TRACKING_STATUS_COLUMN = "追踪状态"
SOURCE_SHEET_COLUMN = "来源sheet"
FROZEN_STATUS = "已冻结"
LEFT_STYLE_END_COLUMN = "返款配置"
PROGRESS_STYLE_START_COLUMN = "打卡数"
TRACKING_GROUP_COLUMN = "追踪分组"
FROZEN_BACKGROUND_COLOR = "#F2F2F2"
FROZEN_TEXT_COLOR = "#6B7280"


@dataclass
class SourceDocument:
    title: str
    columns: list[str]
    rows: list[Any]
    cell_meta: dict[str, Any]
    data_start_row: int
    row_by_key: dict[str, int]


def _normalize_text(value: Any) -> str:
    return str(value or "").strip()


def _parse_datetime(value: Any) -> datetime | None:
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime(value.year, value.month, value.day)
    text = _normalize_text(value)
    if not text:
        return None
    normalized = text.replace("年", "-").replace("月", "-").replace("日", "")
    normalized = normalized.replace("/", "-").replace(".", "-")
    match = re.search(r"(\d{4})-(\d{1,2})-(\d{1,2})(?:\D+(\d{1,2}):(\d{1,2})(?::(\d{1,2}))?)?", normalized)
    if not match:
        return None
    try:
        return datetime(
            int(match.group(1)),
            int(match.group(2)),
            int(match.group(3)),
            int(match.group(4) or 0),
            int(match.group(5) or 0),
            int(match.group(6) or 0),
        )
    except ValueError:
        return None


def _normalize_datetime_key(value: Any) -> str:
    parsed = _parse_datetime(value)
    return parsed.isoformat(sep=" ", timespec="minutes") if parsed is not None else _normalize_text(value)


def _row_value(row: list[Any], columns: list[str], header: str) -> Any:
    if header not in columns:
        return ""
    index = columns.index(header)
    return row[index] if index < len(row) else ""


def _row_key(row: list[Any], columns: list[str]) -> str:
    return "|".join([
        _normalize_datetime_key(_row_value(row, columns, "报名日期")),
        _normalize_text(_row_value(row, columns, "姓名")),
        _normalize_text(_row_value(row, columns, "商户订单号")),
        _normalize_text(_row_value(row, columns, "用户ID")),
    ])


def _load_source_documents(workbook_path: Path) -> dict[str, SourceDocument]:
    workbook = load_workbook(workbook_path, data_only=False)
    result: dict[str, SourceDocument] = {}
    for title in SOURCE_SHEET_TITLES:
        document = _attendance_document(workbook[title])
        columns = list(document.get("columns") or [])
        rows = list(document.get("rows") or [])
        row_by_key: dict[str, int] = {}
        for row_index, row in enumerate(rows):
            normalized_row = list(row)
            row_by_key.setdefault(_row_key(normalized_row, columns), row_index)
        result[title] = SourceDocument(
            title=title,
            columns=columns,
            rows=rows,
            cell_meta=copy.deepcopy(document.get("cell_meta") or {}),
            data_start_row=int(document.get("data_start_row") or 0),
            row_by_key=row_by_key,
        )
    return result


def _strip_frozen_style(meta: dict[str, Any] | None) -> dict[str, Any] | None:
    if not isinstance(meta, dict):
        return None
    next_meta = copy.deepcopy(meta)
    style = dict(next_meta.get("style") or {})
    if style.get("background_color") == FROZEN_BACKGROUND_COLOR:
        style.pop("background_color", None)
    if style.get("text_color") == FROZEN_TEXT_COLOR:
        style.pop("text_color", None)
    if style:
        next_meta["style"] = style
    else:
        next_meta.pop("style", None)
    return next_meta if next_meta else None


def _set_left_frozen_style(cell_meta: dict[str, Any], key: str) -> None:
    current = dict(cell_meta.get(key) or {})
    current["style"] = {
        **dict(current.get("style") or {}),
        "background_color": FROZEN_BACKGROUND_COLOR,
        "text_color": FROZEN_TEXT_COLOR,
    }
    cell_meta[key] = current


def _replace_or_delete_meta(cell_meta: dict[str, Any], key: str, meta: dict[str, Any] | None) -> bool:
    current = cell_meta.get(key)
    if meta:
        if current != meta:
            cell_meta[key] = meta
            return True
        return False
    if key in cell_meta:
        del cell_meta[key]
        return True
    return False


def run(*, apply: bool, workbook_path: Path = DEFAULT_WORKBOOK_PATH) -> dict[str, Any]:
    source_documents = _load_source_documents(workbook_path)
    with Session(engine) as session:
        sheet = session.exec(select(SheetDocument).where(SheetDocument.numeric_id == ATTENDANCE_SHEET_NUMERIC_ID)).first()
        if sheet is None:
            raise RuntimeError(f"找不到考勤表 numeric_id={ATTENDANCE_SHEET_NUMERIC_ID}")
        document = copy.deepcopy(dict(sheet.document_json or {}))
        columns = _normalize_document_columns(document)
        rows = _extract_document_rows(document)
        data_start_row = _normalize_document_data_start_row(document)
        cell_meta = copy.deepcopy(document.get("cell_meta") or {})
        status_index = columns.index(TRACKING_STATUS_COLUMN)
        source_title_index = columns.index(SOURCE_SHEET_COLUMN)
        left_end = columns.index(LEFT_STYLE_END_COLUMN)
        progress_start = columns.index(PROGRESS_STYLE_START_COLUMN)
        progress_end = columns.index(TRACKING_GROUP_COLUMN) if TRACKING_GROUP_COLUMN in columns else len(columns)

        matched = 0
        unmatched = 0
        changed_cells = 0
        restored_progress_cells = 0
        removed_right_gray_cells = 0
        left_gray_cells = 0
        rule_based_style_cells = 0

        for row_index, row in enumerate(rows):
            row_values = list(row)
            document_row = data_start_row + row_index
            frozen = _normalize_text(row_values[status_index]) == FROZEN_STATUS
            if frozen:
                for column_index in range(left_end + 1):
                    key = f"{document_row}:{column_index}"
                    before = copy.deepcopy(cell_meta.get(key))
                    _set_left_frozen_style(cell_meta, key)
                    if before != cell_meta.get(key):
                        changed_cells += 1
                    left_gray_cells += 1

            source_title = _normalize_text(row_values[source_title_index])
            source = source_documents.get(source_title)
            source_row_index = source.row_by_key.get(_row_key(row_values, columns)) if source else None
            if source is None or source_row_index is None:
                unmatched += 1
                for column_index in range(left_end + 1, len(columns)):
                    key = f"{document_row}:{column_index}"
                    cleaned = _strip_frozen_style(cell_meta.get(key))
                    if _replace_or_delete_meta(cell_meta, key, cleaned):
                        changed_cells += 1
                        removed_right_gray_cells += 1
                continue

            matched += 1
            for column_index in range(progress_start, progress_end):
                source_column = columns[column_index]
                key = f"{document_row}:{column_index}"
                if source_column in source.columns:
                    source_column_index = source.columns.index(source_column)
                    source_meta = copy.deepcopy(
                        source.cell_meta.get(f"{source.data_start_row + source_row_index}:{source_column_index}")
                    )
                else:
                    source_meta = None
                if _replace_or_delete_meta(cell_meta, key, source_meta):
                    changed_cells += 1
                    restored_progress_cells += 1

            for column_index in [*range(left_end + 1, progress_start), *range(progress_end, len(columns))]:
                key = f"{document_row}:{column_index}"
                cleaned = _strip_frozen_style(cell_meta.get(key))
                if _replace_or_delete_meta(cell_meta, key, cleaned):
                    changed_cells += 1
                    removed_right_gray_cells += 1

        styled_document, style_summary = apply_nianzhu_course_progress_styles_to_document({
            **document,
            "cell_meta": cell_meta,
        })
        cell_meta = copy.deepcopy(styled_document.get("cell_meta") or {})
        rule_based_style_cells = int(style_summary.get("changed_cells") or 0)
        changed_cells += rule_based_style_cells

        summary = {
            "rows": len(rows),
            "matched": matched,
            "unmatched": unmatched,
            "changed_cells": changed_cells,
            "left_gray_cells": left_gray_cells,
            "restored_progress_cells": restored_progress_cells,
            "removed_right_gray_cells": removed_right_gray_cells,
            "rule_based_style_cells": rule_based_style_cells,
        }
        if apply and changed_cells:
            document["cell_meta"] = cell_meta
            sheet.document_json = document
            sheet.version = int(sheet.version or 1) + 1
            sheet.updated_at = time.time()
            session.add(sheet)
            session.commit()
        else:
            session.rollback()
        return summary


def main() -> None:
    parser = argparse.ArgumentParser(description="恢复念住闯关考勤表右侧课程进度样式，仅保留左侧冻结灰底。")
    parser.add_argument("--apply", action="store_true", help="实际写入数据库。默认只 dry-run。")
    parser.add_argument("--workbook", default=str(DEFAULT_WORKBOOK_PATH), help="原始 xlsx 路径。")
    args = parser.parse_args()

    summary = run(apply=args.apply, workbook_path=Path(args.workbook))
    mode = "APPLY" if args.apply else "DRY-RUN"
    print(f"[{mode}] 念住闯关考勤表样式修正")
    for key, value in summary.items():
        print(f"{key}: {value}")


if __name__ == "__main__":
    main()
