"""Read uploaded workbooks, preserving dates when an export has empty fill XML."""
import io
import xml.etree.ElementTree as ET
from zipfile import ZipFile


def load_import_workbook(raw_bytes: bytes):
    """Open an XLSX for values; repair only malformed empty style fills in memory.

    Some online-sheet exports contain ``<fill/>``. openpyxl rejects these
    before reading cells. Replacing that empty fill with a no-pattern fill
    preserves number formats, dates, sheet content and the original upload.
    Other corruption remains an explicit read failure.
    """
    from openpyxl import load_workbook
    try:
        return load_workbook(io.BytesIO(raw_bytes), read_only=True, data_only=True)
    except TypeError as exc:
        if "Fill" not in str(exc):
            raise
        ns = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
        with ZipFile(io.BytesIO(raw_bytes)) as source:
            styles = ET.fromstring(source.read('xl/styles.xml'))
            empty = [fill for fill in styles.findall(f'{ns}fills/{ns}fill') if len(fill) == 0]
            if not empty:
                raise
            for fill in empty:
                ET.SubElement(fill, f'{ns}patternFill')
            repaired = io.BytesIO()
            with ZipFile(repaired, 'w') as target:
                for entry in source.infolist():
                    target.writestr(entry, ET.tostring(styles) if entry.filename == 'xl/styles.xml'
                                    else source.read(entry.filename))
        repaired.seek(0)
        return load_workbook(repaired, read_only=True, data_only=True)
